"""Keep PDF pagination, artwork and editable text at their original coordinates."""
from pathlib import Path
from copy import deepcopy
import sys, uuid, json, struct
import pymupdf as fitz
from lxml import etree
from docx import Document
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.part import Part
from docx.opc.packuri import PackURI
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT=Path(__file__).resolve().parent.parent
SOURCE=ROOT/'output/pdf/AutoDrive_Research_Paper_Reference_Style.pdf'
TMP=ROOT/'tmp/docx_pdf_layout';TMP.mkdir(parents=True,exist_ok=True)
LIMIT=int(sys.argv[1]) if len(sys.argv)>1 else 0
OUT=(TMP/'prototype.docx') if LIMIT else ROOT/'output/docx/AutoDrive_Research_Paper_PDF_Layout.docx'
pdf=fitz.open(SOURCE); doc=Document();section=doc.sections[0]
W,H=pdf[0].rect.width,pdf[0].rect.height
section.page_width=Pt(W);section.page_height=Pt(H)
section.top_margin=section.bottom_margin=section.left_margin=section.right_margin=Pt(0)
section.header_distance=section.footer_distance=Pt(0)
normal=doc.styles['Normal'];normal.font.name='Cmr10';normal.font.size=Pt(1)
normal.paragraph_format.space_before=normal.paragraph_format.space_after=Pt(0)
normal.paragraph_format.line_spacing=Pt(1)
for style in doc.styles:
 for n in list(style.element.iter(qn('w:pBdr'))):n.getparent().remove(n)
for name in ['embedTrueTypeFonts','saveSubsetFonts']:
 e=OxmlElement('w:'+name);e.set(qn('w:val'),'true' if name=='embedTrueTypeFonts' else 'false');doc.settings.element.append(e)

fontpart=doc.part.part_related_by(RT.FONT_TABLE)
fontxml=etree.fromstring(fontpart.blob)
for name,file in [('Cmr10','cmr10.ttf'),('Cmb10','cmb10.ttf'),('Cmti10','cmti10.ttf')]:
 raw=bytearray((ROOT/'research_paper/fonts'/file).read_bytes())
 key=uuid.uuid4();mask=key.bytes[::-1]
 for i in range(32):raw[i]^=mask[i%16]
 part=Part(PackURI('/word/fonts/'+name+'.odttf'),'application/vnd.openxmlformats-officedocument.obfuscatedFont',bytes(raw),doc.part.package)
 rid=fontpart.relate_to(part,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/font')
 f=OxmlElement('w:font');f.set(qn('w:name'),name)
 fam=OxmlElement('w:family');fam.set(qn('w:val'),'roman');f.append(fam)
 emb=OxmlElement('w:embedRegular');emb.set(qn('r:id'),rid);emb.set(qn('w:fontKey'),'{'+str(key).upper()+'}');f.append(emb)
 fontxml.append(f)
fontpart._blob=etree.tostring(fontxml,xml_declaration=True,encoding='UTF-8',standalone=True)

V='urn:schemas-microsoft-com:vml';O='urn:schemas-microsoft-com:office:office'
def ve(tag):return etree.Element('{'+V+'}'+tag)
def floating_shape(anchor,x,y,w,h,shapeid,background=False):
 pict=OxmlElement('w:pict');shape=ve('shape')
 shape.set('id',shapeid);shape.set('type','#_x0000_t75' if background else '#_x0000_t202')
 shape.set('style',f'position:absolute;margin-left:{x:.3f}pt;margin-top:{y:.3f}pt;width:{w:.3f}pt;height:{h:.3f}pt;z-index:{-251658240 if background else 251659264};mso-position-horizontal-relative:page;mso-position-vertical-relative:page')
 shape.set('stroked','f');shape.set('filled','f');pict.append(shape)
 anchor.add_run()._r.append(pict)
 return shape

def make_line(anchor,line,pageindex,index):
 original=line['spans']
 origins=[s['origin'][1] for s in original]
 if len(original)>1 and max(origins)-min(origins)>.2:
  # Word clips raised/lowered runs inside an exact-height line. Give each
  # mathematical span its recorded baseline instead.
  for k,span in enumerate(original):
   make_line(anchor,{'spans':[span],'bbox':span['bbox']},pageindex,str(index)+'_'+str(k))
  return
 spans=[]
 for s in line['spans']:
  s=dict(s);s['text']=''.join(c['c'] for c in s['chars']);spans.append(s)
 if not spans:return
 x0,y0,x1,y1=line['bbox'];baseline=max(s['origin'][1] for s in spans)
 visible=[c for s in spans for c in s['chars'] if not c['c'].isspace()]
 if not visible:return
 x0=visible[0]['origin'][0]
 href=next((entry.get('uri') for entry in current_links if entry.get('uri') and fitz.Rect(entry['from']).intersects(fitz.Rect(line['bbox'])) and (fitz.Rect(entry['from']) & fitz.Rect(line['bbox'])).width > (x1-x0)*.5),None)
 fs=max(s['size'] for s in spans)
 # Word's exact-line text frame baseline is calibrated against the native render.
 top=baseline-fs*.963
 shape=floating_shape(anchor,x0,top,max(x1-x0+8,1),fs*1.7,f'text_{pageindex}_{index}')
 box=ve('textbox');box.set('inset','0,0,0,0');box.set('style','mso-fit-shape-to-text:f')
 content=OxmlElement('w:txbxContent');p=OxmlElement('w:p');pr=OxmlElement('w:pPr')
 spacing=OxmlElement('w:spacing');spacing.set(qn('w:before'),'0');spacing.set(qn('w:after'),'0');spacing.set(qn('w:line'),str(round(fs*1.2*20)));spacing.set(qn('w:lineRule'),'exact');pr.append(spacing)
 jc=OxmlElement('w:jc');jc.set(qn('w:val'),'left');pr.append(jc)
 tabs=OxmlElement('w:tabs');pr.append(tabs);p.append(pr)
 tokens=[];word=[]
 for s in spans:
  for char in s['chars']:
   if char['c'].isspace():
    if word:tokens.append(word);word=[]
   else:word.append((char,s))
 if word:tokens.append(word)
 first=True
 for word in tokens:
  if not first:
   stop=OxmlElement('w:tab');stop.set(qn('w:val'),'left');stop.set(qn('w:pos'),str(round((word[0][0]['origin'][0]-x0)*20)));tabs.append(stop)
   tr=OxmlElement('w:r');tr.append(OxmlElement('w:tab'));p.append(tr)
  first=False
  groups=[]
  for char,s in word:
   if groups and groups[-1][0] is s:groups[-1][1].append(char['c'])
   else:groups.append([s,[char['c']]])
  for s,letters in groups:
   add_run(p,s,''.join(letters),baseline,href)
 content.append(p);box.append(content);shape.append(box)

def add_run(p,s,text,baseline,href=None):
  r=OxmlElement('w:r');rp=OxmlElement('w:rPr')
  fonts=OxmlElement('w:rFonts');font={'Helvetica':'Arial','ArialMT':'Arial','TimesNewRomanPSMT':'Times New Roman'}.get(s['font'],s['font'])
  for k in ['ascii','hAnsi','cs','eastAsia']:fonts.set(qn('w:'+k),font)
  rp.append(fonts)
  size=OxmlElement('w:sz');size.set(qn('w:val'),str(round(s['size']*2)));rp.append(size)
  color=OxmlElement('w:color');color.set(qn('w:val'),f'{s["color"]:06X}');rp.append(color)
  rp.append(OxmlElement('w:noProof'))
  # Preserve superscript/subscript placement recorded in the PDF spans.
  dy=baseline-s['origin'][1]
  if abs(dy)>.2:
   pos=OxmlElement('w:position');pos.set(qn('w:val'),str(round(dy*2)));rp.append(pos)
  r.append(rp);t=OxmlElement('w:t');t.set(qn('xml:space'),'preserve');t.text=text;r.append(t)
  if href:
   link=OxmlElement('w:hyperlink');link.set(qn('r:id'),doc.part.relate_to(href,RT.HYPERLINK,is_external=True));link.append(r);p.append(link)
  else:p.append(r)

pages=list(pdf)[:LIMIT or len(pdf)];count=0
for i,page in enumerate(pages):
 current_links=page.get_links()
 anchor=doc.add_paragraph()
 anchor.paragraph_format.line_spacing=Pt(1)
 if i:anchor.paragraph_format.page_break_before=True
 # Keep the full-colour vector/image artwork; all PDF text becomes native Word text.
 artwork=fitz.open();artwork.insert_pdf(pdf,from_page=i,to_page=i);ap=artwork[0]
 for b in ap.get_text('dict')['blocks']:
  if 'lines' in b:
   for l in b['lines']:
    for s in l['spans']:ap.add_redact_annot(s['bbox'],fill=False)
 ap.apply_redactions(images=0,graphics=0,text=0)
 path=TMP/f'artwork-{i+1}.png';ap.get_pixmap(matrix=fitz.Matrix(3,3),alpha=False).save(path)
 rid,_=doc.part.get_or_add_image(str(path))
 bg=floating_shape(anchor,0,0,W,H,f'artwork_{i+1}',True)
 im=ve('imagedata');im.set(qn('r:id'),rid);bg.append(im)
 for b in page.get_text('rawdict')['blocks']:
  if 'lines' in b:
   for l in b['lines']:
    make_line(anchor,l,i,count);count+=1
doc.core_properties.title='AutoDrive: Real-Time Autonomous Driving System Using Computer Vision and Deep Learning'
doc.core_properties.author='Harsh Bhardwaj; Shrishti Singh; Rajat Kandpal'
doc.core_properties.subject='PDF layout preserved page by page with editable text'
doc.save(OUT)
print(json.dumps({'output':str(OUT),'pages':len(pages),'editable_lines':count},indent=2))
