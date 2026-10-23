"""Export the current manuscript flowables into an editable Word document."""
from pathlib import Path
from io import BytesIO
from html.parser import HTMLParser
import sys, json
import pymupdf as fitz
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from reportlab.platypus import Paragraph, Table, Image, KeepTogether, Spacer
from reportlab.graphics.shapes import Drawing
from reportlab.graphics import renderPDF

ROOT=Path(__file__).resolve().parent
sys.path.insert(0,str(ROOT))
ns={'__file__':str(ROOT/'build_reference_style.py'),'__name__':'manuscript_source'}
code=(ROOT/'build_reference_style.py').read_text(encoding='utf-8').split('doc=Article(); doc.build(S)')[0]
exec(compile(code,str(ROOT/'build_reference_style.py'),'exec'),ns)
OUT=ROOT.parent/'output/docx/AutoDrive_Research_Paper_Latest.docx'
TMP=ROOT.parent/'tmp/docx_latest'
OUT.parent.mkdir(parents=True,exist_ok=True)
TMP.mkdir(parents=True,exist_ok=True)
doc=Document()
sec=doc.sections[0]
sec.page_width=Pt(ns['W']); sec.page_height=Pt(ns['H'])
sec.left_margin=Pt(100.1); sec.right_margin=Pt(124.55)
sec.top_margin=Pt(78); sec.bottom_margin=Pt(56)
sec.header_distance=Pt(20); sec.footer_distance=Pt(22)
mir=OxmlElement('w:mirrorMargins'); doc.settings.element.append(mir)
normal=doc.styles['Normal']
normal.font.name='Times New Roman'; normal.font.size=Pt(9.96)
normal.font.color.rgb=__import__('docx').shared.RGBColor(0,0,0)
for name,s in ns['ST'].items():
 target={'Section':'Heading 1','Subsection':'Heading 2'}.get(name,name)
 if target not in doc.styles: doc.styles.add_style(target,1)
 st=doc.styles[target]; st.font.name='Times New Roman'; st.font.size=Pt(s.fontSize)
 st.font.bold='Bold' in s.fontName; st.font.italic='Italic' in s.fontName
 st.font.color.rgb=__import__('docx').shared.RGBColor(0,0,0)
 pf=st.paragraph_format
 pf.line_spacing=Pt(s.leading); pf.space_before=Pt(s.spaceBefore); pf.space_after=Pt(s.spaceAfter)
 pf.left_indent=Pt(s.leftIndent); pf.first_line_indent=Pt(s.firstLineIndent)
 pf.widow_control=True; pf.keep_with_next=False; pf.keep_together=False
 pf.alignment={0:WD_ALIGN_PARAGRAPH.LEFT,1:WD_ALIGN_PARAGRAPH.CENTER,2:WD_ALIGN_PARAGRAPH.RIGHT,4:WD_ALIGN_PARAGRAPH.JUSTIFY}.get(s.alignment,WD_ALIGN_PARAGRAPH.LEFT)
 if name in ['Section','Subsection','TableCaption','AbstractLabel']: pf.keep_with_next=True

class RichText(HTMLParser):
 def __init__(self,p):
  super().__init__(convert_charrefs=True); self.p=p; self.stack=[{}]
 def handle_starttag(self,tag,attrs):
  a=dict(attrs)
  if tag=='br': self.p.add_run().add_break(); return
  cur=self.stack[-1].copy()
  if tag in ['b','strong']: cur['bold']=True
  if tag in ['i','em']: cur['italic']=True
  if tag=='sub': cur['subscript']=True
  if tag in ['super','sup']: cur['superscript']=True
  if tag in ['link','a']: cur['href']=a.get('href')
  self.stack.append(cur)
 def handle_endtag(self,tag):
  if len(self.stack)>1 and tag!='br': self.stack.pop()
 def handle_data(self,data):
  if not data:return
  fmt=self.stack[-1]; r=self.p.add_run(data)
  for k in ['bold','italic','subscript','superscript']:
   if k in fmt:setattr(r.font,k,fmt[k])
  if fmt.get('href'):
   link=OxmlElement('w:hyperlink'); link.set(qn('r:id'),self.p.part.relate_to(fmt['href'],RT.HYPERLINK,is_external=True))
   r.font.color.rgb=__import__('docx').shared.RGBColor(0,0,0)
   self.p._p.remove(r._r); link.append(r._r); self.p._p.append(link)

def paragraph(flow,parent=doc,existing=None):
 sty={'Section':'Heading 1','Subsection':'Heading 2'}.get(flow.style.name,flow.style.name)
 p=existing if existing is not None else parent.add_paragraph()
 p.style=doc.styles[sty]
 RichText(p).feed(flow.text)
 if flow.style.name in ['Abstract','AbstractLabel']:
  p.paragraph_format.left_indent=Pt(24);p.paragraph_format.right_indent=Pt(24)
 if flow.style.name=='Title':p.paragraph_format.space_before=Pt(54)
 return p

def raster_pdf(pdf,key):
 path=TMP/(key+'.png')
 pdf[0].get_pixmap(matrix=fitz.Matrix(4,4),alpha=False).save(path)
 return path

def picture(flow,parent=doc):
 p=parent.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
 pf=p.paragraph_format;pf.space_before=Pt(3);pf.space_after=Pt(0);pf.line_spacing=1
 if isinstance(flow,Image): path=flow.filename;w,h=flow.drawWidth,flow.drawHeight
 elif isinstance(flow,Drawing):
  pdf=fitz.open(stream=renderPDF.drawToString(flow),filetype='pdf')
  path=raster_pdf(pdf,'steering_readout_diagram');w,h=flow.width,flow.height
 else:
  path=raster_pdf(flow.pdf,flow.key);w,h=flow.width,flow.height
 p.add_run().add_picture(str(path),width=Pt(w),height=Pt(h))
 p.paragraph_format.keep_with_next=True
 return p

def border(cell,edge,val='single',sz='4'):
 tcPr=cell._tc.get_or_add_tcPr(); b=tcPr.find(qn('w:tcBorders'))
 if b is None:b=OxmlElement('w:tcBorders');tcPr.append(b)
 el=OxmlElement('w:'+edge);el.set(qn('w:val'),val);el.set(qn('w:sz'),sz);el.set(qn('w:color'),'000000');b.append(el)

def table(flow,parent=doc):
 cells=flow._cellvalues
 # The source abstract uses a one-column layout wrapper, represented as prose.
 if len(cells[0])==1 and isinstance(cells[0][0],Paragraph) and cells[0][0].style.name=='AbstractLabel':
  for row in cells:
   for cell in row:paragraph(cell,parent)
  return
 # Equation layout: preserve editable mathematical text and its number.
 if len(cells)==1 and all(isinstance(c,Paragraph) and c.style.name=='Equation' for c in cells[0]):
  p=paragraph(cells[0][0],parent);p.add_run('   '+cells[0][1].getPlainText());return
 t=parent.add_table(rows=len(cells),cols=len(cells[0])); t.autofit=False
 t.alignment=__import__('docx').enum.table.WD_TABLE_ALIGNMENT.CENTER
 widths=flow._colWidths
 for j,width in enumerate(widths):t.columns[j].width=Pt(width)
 is_images=all(isinstance(c,Image) for row in cells for c in row)
 for i,row in enumerate(cells):
  for j,value in enumerate(row):
   cell=t.cell(i,j);cell.width=Pt(widths[j]);cell.vertical_alignment=__import__('docx').enum.table.WD_CELL_VERTICAL_ALIGNMENT.CENTER
   pr=cell._tc.get_or_add_tcPr();margin=OxmlElement('w:tcMar')
   for side,v in [('top',40),('bottom',40),('left',40),('right',40)]:
    e=OxmlElement('w:'+side);e.set(qn('w:w'),str(0 if is_images else v));e.set(qn('w:type'),'dxa');margin.append(e)
   pr.append(margin)
   for edge in ['top','bottom','left','right','insideH','insideV']:border(cell,edge,'nil')
   if not is_images:
    if i==0:border(cell,'top',sz='6');border(cell,'bottom',sz='4')
    if i==len(cells)-1:border(cell,'bottom',sz='6')
   vals=value if isinstance(value,(list,tuple)) else [value]
   first=True
   for v in vals:
    if isinstance(v,Paragraph):paragraph(v,cell,cell.paragraphs[0] if first else None)
    elif isinstance(v,Image):
     p=cell.paragraphs[0] if first else cell.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
     p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1
     p.add_run().add_picture(v.filename,width=Pt(v.drawWidth),height=Pt(v.drawHeight))
    else:cell.paragraphs[0].text=str(v)
    first=False
  trPr=t.rows[i]._tr.get_or_add_trPr()
  dont=OxmlElement('w:cantSplit');trPr.append(dont)
  if i==0 and not is_images:
   repeat=OxmlElement('w:tblHeader');trPr.append(repeat)
  if is_images:
   for cell in t.rows[i].cells:
    for p in cell.paragraphs:p.paragraph_format.keep_with_next=True

def convert(flow):
 if isinstance(flow,KeepTogether):
  for f in flow._content:convert(f)
 elif isinstance(flow,Paragraph):paragraph(flow)
 elif isinstance(flow,Table):table(flow)
 elif isinstance(flow,(Image,Drawing)) or hasattr(flow,'pdf'):picture(flow)
 elif isinstance(flow,Spacer):pass

for f in ns['S']:convert(f)
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
footer.paragraph_format.space_after=Pt(0);footer.paragraph_format.line_spacing=Pt(11)
run=footer.add_run();run.font.name='Times New Roman';run.font.size=Pt(9.96)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');run._r.addnext(field)
doc.core_properties.title='AutoDrive: Real-Time Autonomous Driving System Using Computer Vision and Deep Learning'
doc.core_properties.author='Harsh Bhardwaj; Shrishti Singh; Rajat Kandpal'
doc.core_properties.subject='Latest research manuscript with recorded perception and steering experiments'
for style in doc.styles:
 for node in list(style.element.iter(qn('w:pBdr'))):node.getparent().remove(node)
 for fonts in style.element.iter(qn('w:rFonts')):
  for attr in list(fonts.attrib):
   if attr.endswith('Theme'):del fonts.attrib[attr]
for node in list(doc.element.iter(qn('w:pBdr'))):node.getparent().remove(node)
doc.save(OUT)
print(json.dumps({'output':str(OUT),'paragraphs':len(doc.paragraphs),'tables':len(doc.tables),'figures':len(doc.inline_shapes)},indent=2))
